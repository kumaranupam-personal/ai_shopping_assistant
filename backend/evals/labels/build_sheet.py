"""Builds the labeling sheet for calibrating the LLM judge (docs/07-evaluation.md, Labeling guide; docs/08-roadmap.md, Phase 19).

From the dev runs that store tool calls, it picks 45 real turns and 15 deliberately broken copies, shuffles them and writes:
- `labeling-sheet.xlsx`: the rows to label, showing only what the rubric grades;
- `recheck-sheet.xlsx`: a fifth of the rows again, reshuffled, for the labeler's consistency check after a break;
- `key/key.json`: what the sheets hide (source run, provider, case, rule-check result, the break, the calibration split).

The breaks themselves are written by hand in `key/broken.json`. Run from `backend/` (openpyxl isn't a project dependency):
`uv run --with openpyxl python -m evals.labels.build_sheet` (add `--sources` to print the turns chosen for breaking).
"""

import json
import random
import re
import sys
from pathlib import Path

from app.config import Settings
from app.search.engine import _summary
from app.search.index import load_index
from evals.run import RESULTS_DIR, catalog_fingerprint
from evals.turn_view import INFO, render

LABELS_DIR = Path(__file__).resolve().parent
KEY_DIR = LABELS_DIR / "key"
RUNS = sorted(RESULTS_DIR.glob("dev/20261013-23*-dev-*-a8130bf25dd1.json"))  # the dev runs with tool calls
SEED = 20261014
CRITERIA = ["Relevance", "Helpfulness", "Faithfulness", "Relaxation honesty", "Clarification quality", "Decline quality", "Language and style"]
# How many broken copies break each criterion, and which turns can carry that break.
BREAKS = {
    "Relevance": (2, {"search", "refinement"}),
    "Helpfulness": (2, {"search", "refinement"}),
    "Faithfulness": (3, {"search", "refinement", "question"}),
    "Relaxation honesty": (2, {"relax"}),
    "Clarification quality": (1, {"clarify"}),
    "Decline quality": (1, {"decline"}),
    "Language and style": (4, {"search", "refinement", "question"}),
}


def behavior(case: dict, turn: int) -> str:
    keys = case["expect"]
    if turn < len(case["turns"]) - 1:
        return "search"  # an earlier turn of a multi-turn case
    if "clarifies" in keys:
        return "clarify"
    if "declines" in keys:
        return "decline"
    if "relax" in case["id"]:
        return "relax"
    if "mentions" in keys:
        return "question"
    return "refinement" if len(case["turns"]) > 1 else "search"


def load_turns() -> list[dict]:
    from evals.run import SETS_DIR, load_cases
    cases = {c.id: c.model_dump(exclude_unset=True) for c in load_cases(SETS_DIR / "dev.yaml", [])}
    turns = []
    for run in RUNS:
        data = json.loads(run.read_text())
        for case in data["cases"]:
            for i, turn in enumerate(case["turns"]):
                turns.append({
                    "run": run.name, "provider": data["provider"], "catalog_sha256": data["catalog_sha256"],
                    "case": case["id"], "turn": i, "behavior": behavior(cases[case["id"]], i),
                    "hinglish": case["id"].startswith("hinglish"), "case_passed": case["passed"],
                    "history": case["turns"][:i], "data": turn,
                })
    return turns


def select(turns: list[dict]) -> tuple[list[dict], list[dict]]:
    """15 broken sources from distinct turns, then 45 real rows from the rest, 7 of them a second provider's reply."""
    rng = random.Random(SEED)
    by_key: dict[tuple, list[dict]] = {}
    for t in turns:
        by_key.setdefault((t["case"], t["turn"]), []).append(t)
    keys = sorted(by_key)
    used, broken = set(), []
    for criterion, (count, behaviors) in BREAKS.items():
        eligible = [k for k in keys if k not in used and by_key[k][0]["behavior"] in behaviors]
        for k in rng.sample(eligible, count):
            used.add(k)
            broken.append({**rng.choice(by_key[k]), "criterion": criterion})
    rest = [k for k in keys if k not in used]
    repeats = rng.sample(rest, 45 - len(rest))  # these turns appear twice, with two providers' replies
    # Providers for the real rows, so that with the broken ones each provider has 20 rows.
    providers = sorted({t["provider"] for t in turns})
    need = [p for p in providers for _ in range(20 - sum(b["provider"] == p for b in broken))]
    balance = random.Random(SEED + 4)
    while True:
        balance.shuffle(need)
        firsts, seconds = need[:len(rest)], need[len(rest):]
        if all(seconds[i] != firsts[rest.index(k)] for i, k in enumerate(repeats)):
            break
    pick = {(t["case"], t["turn"], t["provider"]): t for t in turns}
    real = [pick[(*k, p)] for k, p in zip(rest, firsts)] + [pick[(*k, p)] for k, p in zip(repeats, seconds)]
    return real, broken


WIDTHS = {"Row": 7, "Earlier conversation": 40, "User message": 28, "What the agent did": 48, "Products shown": 70,
          "Headline and suggestions": 30, "Reply": 48, "Length": 12, "Notes on fails": 36, "Other comments": 30}


def write_sheet(path: Path, rows: list[tuple[str, dict]], rubric_rows: list[list[str]], title: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    header_fill, input_fill = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFF2CC")
    thin = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
    wb = Workbook()

    guide = wb.active
    guide.title = "Instructions"
    lines = [
        (title, bold),
        ("", font),
        ("What to do", bold),
        ("1. Read the 'Rubric' tab once, especially decisions D1 to D8. Label against it, not a personal standard.", font),
        ("2. In the 'Labels' tab, read a row's grey columns left to right: earlier conversation, user message, what the agent did, the products shown, then the reply.", font),
        ("3. Fill the yellow cells: pass, fail or n/a for each criterion (pick from the dropdown), a note for every fail saying which criterion and why, and any other comment.", font),
        ("4. Use n/a only when the criterion doesn't apply to the turn (see 'Applies when' in the Rubric tab).", font),
        ("5. Check facts in the reply against 'What the agent did' and 'Products shown', not against memory.", font),
        ("6. Label in two or three sittings. Don't open the key folder (evals/labels/key/) until every row is labeled.", font),
        ("", font),
        ("Example row (made up, not one of the rows to label)", bold),
        ("User: 'Steel water bottle under ₹800'. Agent searched kitchen items ≤ ₹800 and showed 4 bottles. Reply: 'The first keeps drinks cold longest, and the third is the lightest to carry.'", font),
        ("Relevance: pass | Helpfulness: pass | Faithfulness: fail | Relaxation honesty: n/a | Clarification quality: n/a | Decline quality: n/a | Language and style: pass", font),
        ("Notes on fails: 'Faithfulness: nothing in the products shown says how long the first keeps drinks cold.'", font),
    ]
    for r, (text, f) in enumerate(lines, 1):
        cell = guide.cell(row=r, column=1, value=text)
        cell.font, cell.alignment = f, wrap
    guide.column_dimensions["A"].width = 140

    rubric = wb.create_sheet("Rubric")
    for r, values in enumerate(rubric_rows, 1):
        for c, value in enumerate(values, 1):
            cell = rubric.cell(row=r, column=c, value=value)
            cell.font, cell.alignment = (bold if r == 1 or len(values) == 1 and value.startswith("Decision") else font), wrap
    for col, width in zip("ABCDE", (22, 30, 45, 45, 16)):
        rubric.column_dimensions[col].width = width

    labels = wb.create_sheet("Labels")
    headers = ["Row", *INFO, *CRITERIA, "Notes on fails", "Other comments"]
    for c, name in enumerate(headers, 1):
        cell = labels.cell(row=1, column=c, value=name)
        cell.font, cell.alignment, cell.fill, cell.border = bold, wrap, header_fill, thin
        labels.column_dimensions[cell.column_letter].width = WIDTHS.get(name, 14)
    choice = DataValidation(type="list", formula1='"pass,fail,n/a"', allow_blank=True)
    labels.add_data_validation(choice)
    for r, (row_id, fields) in enumerate(rows, 2):
        values = [row_id, *(fields[k] for k in INFO), *([None] * len(CRITERIA)), None, None]
        for c, value in enumerate(values, 1):
            cell = labels.cell(row=r, column=c, value=value)
            cell.font, cell.alignment, cell.border = font, wrap, thin
            if c > 1 + len(INFO):
                cell.fill = input_fill
        first, last = labels.cell(row=r, column=2 + len(INFO)), labels.cell(row=r, column=1 + len(INFO) + len(CRITERIA))
        choice.add(f"{first.coordinate}:{last.coordinate}")
    labels.freeze_panes = "C2"
    wb.active = wb.index(labels)
    wb.save(path)


def has_labels(path: Path) -> bool:
    """Whether any judgment cell of a sheet's Labels tab is filled."""
    if not path.exists():
        return False
    from openpyxl import load_workbook
    sheet = load_workbook(path, read_only=True)["Labels"]
    return any(v not in (None, "") for row in sheet.iter_rows(min_row=2, min_col=2 + len(INFO), values_only=True) for v in row)


def rubric_rows() -> list[list[str]]:
    """The rubric's table, decisions and guide, copied from docs/07-evaluation.md so the sheet can't drift from it."""
    doc = (LABELS_DIR.parents[2] / "docs" / "07-evaluation.md").read_text()
    section = doc[doc.index("## Judge rubric"):]
    rows = [["Criterion", "Applies when", "Pass", "Fail", "From 04-agent.md"]]
    for line in section.splitlines():
        if line.startswith("| **"):
            rows.append([re.sub(r"\*\*", "", cell).strip() for cell in line.strip("|").split("|")])
    rows.append([""])
    for line in section.splitlines():
        if line.startswith("- **D"):
            name, _, text = line[4:].partition(".** ")
            rows.append([f"Decision {name}"])
            rows.append([re.sub(r"\*\*", "", text)])
    return rows


def build() -> None:
    index = load_index(Settings().data_dir)
    turns = load_turns()
    fingerprint = catalog_fingerprint(index)
    if {t["catalog_sha256"] for t in turns} != {fingerprint}:
        raise SystemExit("The catalog doesn't match the one the runs used, so product rows can't be rebuilt from IDs.")
    rows = {row["id"]: _summary(row) for row in index.conn.execute("SELECT * FROM products")}
    real, broken = select(turns)
    edits = json.loads((KEY_DIR / "broken.json").read_text())
    for item, edit in zip(broken, edits, strict=True):
        assert item["criterion"] == edit["criterion"], (item["criterion"], edit)
        assert set(edit.get("replace_shown", {})) <= set(item["data"]["shown"] or []), edit  # a swap must hit a shown product
    items = [(t, None) for t in real] + list(zip(broken, edits))
    rng = random.Random(SEED + 11)  # row order (+1 was replaced before any labeling, after rows were printed while checking)
    rng.shuffle(items)
    sheet_rows, key_rows = [], []
    for n, (item, edit) in enumerate(items, 1):
        row_id = f"T{n:02d}"
        sheet_rows.append((row_id, render(item, rows, edit)))
        key_rows.append({
            "row": row_id, "source_run": item["run"], "provider": item["provider"], "case": item["case"], "turn": item["turn"] + 1,
            "behavior": item["behavior"], "language": "hinglish" if item["hinglish"] else "english",
            "case_passed_rule_checks": item["case_passed"], "broken": edit is not None,
            **({"broken_id": edit["id"], "broken_criterion": edit["criterion"], "change": edit["change"]} if edit else {}),
        })
    # A random third goes to the held-out labels, a third of the broken rows and of the real ones alike.
    for group in (True, False):
        ids = [k["row"] for k in key_rows if k["broken"] is group]
        held = set(random.Random(SEED + 2).sample(ids, len(ids) // 3))
        for k in key_rows:
            if k["row"] in ids:
                k["split"] = "held-out" if k["row"] in held else "calibration"
    recheck = random.Random(SEED + 3).sample([r for r, _ in sheet_rows], 12)
    recheck_rows = [(f"R{n:02d}", dict(sheet_rows)[source]) for n, source in enumerate(recheck, 1)]

    for name in ("labeling-sheet.xlsx", "recheck-sheet.xlsx"):
        if has_labels(LABELS_DIR / name):
            raise SystemExit(f"{name} already has labels in it; move it away first to rebuild the sheets.")
    rubric = rubric_rows()
    write_sheet(LABELS_DIR / "labeling-sheet.xlsx", sheet_rows, rubric, "Labeling sheet: 60 turns to label against the judge rubric")
    write_sheet(LABELS_DIR / "recheck-sheet.xlsx", recheck_rows, rubric,
                "Recheck sheet: 12 turns to label again, at least a few hours after the main sheet, without looking at it")
    KEY_DIR.mkdir(exist_ok=True)
    (KEY_DIR / "key.json").write_text(json.dumps({
        "seed": SEED, "runs": [r.name for r in RUNS], "catalog_sha256": fingerprint, "rows": key_rows,
        "recheck": {f"R{n:02d}": source for n, source in enumerate(recheck, 1)},
    }, indent=2, ensure_ascii=False) + "\n")
    print(f"{len(sheet_rows)} rows ({sum(k['broken'] for k in key_rows)} broken), {len(recheck_rows)} recheck rows, "
          f"{sum(k['split'] == 'held-out' for k in key_rows)} held out")


if __name__ == "__main__" and "--sources" in sys.argv:
    real, broken = select(load_turns())
    print(len(real), "real,", len(broken), "broken")
    for i, b in enumerate(broken, 1):
        d = b["data"]
        print(f"\n## B{i:02d} {b['criterion']} | {b['provider']} {b['case']} turn {b['turn'] + 1} ({b['behavior']})")
        for h in b["history"]:
            print("   earlier user:", h["user"], "\n   earlier reply:", h["reply"])
        print("   user:", d["user"])
        print("   tools:", [(c["name"], c.get("result", {}).get("total_matches") if isinstance(c.get("result"), dict) else None) for c in d["tool_calls"]])
        print("   shown:", d["shown"])
        print("   reply:", d["reply"])
elif __name__ == "__main__":
    build()
