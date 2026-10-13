"""The LLM judge: grades eval turns against the rubric in docs/07-evaluation.md (Judge rubric), and measures its agreement
with human labels (Labeling guide). Phase 19 in docs/08-roadmap.md.

The judge reads a turn exactly as a labeler does (evals/turn_view.py) and answers through one tool call, so every
provider returns the same structured verdicts. It never sees the agent's system prompt.

From `backend/` (every judging call is a paid LLM call; openpyxl only for the sheets):
- `uv run python -m evals.judge --results <results file>`: judge every turn of a dev run (a holdout run needs `--holdout`
  and a recorded plan, see evals/README.md).
- `uv run --with openpyxl python -m evals.judge --sheet calibration|held-out`: judge the labeling sheet's rows of that split.
- `uv run --with openpyxl python -m evals.judge --agreement <judged sheet file>`: compare a judged sheet with the human labels.
- `uv run --with openpyxl python -m evals.judge --consistency`: compare the recheck sheet with the labeling sheet.
"""

import argparse
import asyncio
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from app.config import Settings
from app.llm.base import LLMProvider, ToolSpec
from app.llm.registry import build_provider
from app.search.engine import _summary
from app.search.index import SearchIndex, load_index
from evals.run import RESULTS_DIR, catalog_fingerprint
from evals.turn_view import INFO, render

DOCS = Path(__file__).resolve().parents[2] / "docs" / "07-evaluation.md"
LABELS_DIR = Path(__file__).resolve().parent / "labels"
CRITERIA = {
    "relevance": "Relevance", "helpfulness": "Helpfulness", "faithfulness": "Faithfulness",
    "relaxation_honesty": "Relaxation honesty", "clarification_quality": "Clarification quality",
    "decline_quality": "Decline quality", "language_and_style": "Language and style",
}
VERDICTS = ("pass", "fail", "n/a")

# Refined during calibration against the calibration labels only (docs/08-roadmap.md, Phase 19, Part 6).
INSTRUCTIONS = """You grade one turn of a conversation between a shopper and an online store's shopping assistant, against the rubric below. You are a strict, consistent grader, not the assistant.

For every criterion, decide pass, fail or n/a exactly as the rubric and its decisions D1 to D8 define them, and give a one-line reason (for n/a, say why it doesn't apply). Grade only the current turn's reply, cards and suggestions; earlier turns are context. Check every product fact in the reply against "What the agent did" and "Products shown" in the turn, never against your own knowledge. Count words and sentences as decision D4 says; the turn's "Length" line already gives the counts.

Record your verdicts by calling record_verdicts once."""


def rubric_text() -> str:
    """The rubric and its decisions, from docs/07-evaluation.md, so the judge and the labelers read the same text."""
    doc = DOCS.read_text()
    return doc[doc.index("## Judge rubric"):doc.index("## Labeling guide")].strip()


SYSTEM = f"{INSTRUCTIONS}\n\n{rubric_text()}"
VERDICT_TOOL = ToolSpec(
    name="record_verdicts",
    description="Record the verdict and a one-line reason for every criterion of the rubric.",
    parameters={
        "type": "object",
        "properties": {
            key: {
                "type": "object",
                "description": name,
                "properties": {
                    "verdict": {"type": "string", "enum": list(VERDICTS)},
                    "reason": {"type": "string", "description": "One line: why, or why the criterion doesn't apply."},
                },
                "required": ["verdict", "reason"],
                "additionalProperties": False,
            }
            for key, name in CRITERIA.items()
        },
        "required": list(CRITERIA),
        "additionalProperties": False,
    },
)


def judge_provider(settings: Settings) -> LLMProvider:
    if not settings.judge_provider:
        raise SystemExit("Set JUDGE_PROVIDER (and optionally JUDGE_MODEL and JUDGE_EFFORT) to choose the judge.")
    return build_provider(settings.model_copy(update={
        "llm_provider": settings.judge_provider, "llm_model": settings.judge_model, "llm_effort": settings.judge_effort,
    }))


def judge_version(provider: LLMProvider, effort: str) -> str:
    """A short hash of everything that decides a verdict: instructions, rubric, tool, model and effort."""
    spec = {"system": SYSTEM, "tool": VERDICT_TOOL.parameters, "provider": provider.name, "model": provider.model, "effort": effort}
    return hashlib.sha256(json.dumps(spec, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:12]


def turn_text(fields: dict) -> str:
    return "\n\n".join(f"## {name}\n{fields[name]}" for name in INFO)


def parse_verdicts(arguments: dict) -> dict | None:
    """The verdicts by criterion name, or None when the call misses a criterion or uses an unknown verdict."""
    verdicts = {}
    for key, name in CRITERIA.items():
        item = arguments.get(key)
        if not isinstance(item, dict) or item.get("verdict") not in VERDICTS or not isinstance(item.get("reason"), str):
            return None
        verdicts[name] = {"verdict": item["verdict"], "reason": item["reason"]}
    return verdicts


async def judge_turn(provider: LLMProvider, fields: dict, attempts: int = 2) -> dict:
    """Grades one rendered turn. A response without a valid record_verdicts call is retried once, then recorded as an error."""
    cost, history = 0.0, [provider.user_message(turn_text(fields))]
    for _ in range(attempts):
        response = await provider.complete(SYSTEM, history, [VERDICT_TOOL])
        cost += provider.prices.cost(response.usage) if provider.prices else 0.0
        call = next((c for c in response.tool_calls if c.name == VERDICT_TOOL.name), None)
        if call and (verdicts := parse_verdicts(call.arguments)):
            return {"verdicts": verdicts, "cost_usd": round(cost, 6)}
    return {"error": "no valid record_verdicts call", "cost_usd": round(cost, 6)}


async def judge_all(provider: LLMProvider, items: dict[str, dict], concurrency: int = 4) -> dict[str, dict]:
    gate = asyncio.Semaphore(concurrency)

    async def one(item_id: str, fields: dict):
        async with gate:
            result = await judge_turn(provider, fields)
            print(f"{item_id}: {'error' if 'error' in result else 'judged'}", file=sys.stderr)
            return item_id, result

    return dict(await asyncio.gather(*(one(i, f) for i, f in items.items())))


def summarize_verdicts(judged: dict[str, dict]) -> dict:
    """Per criterion, how many turns it applied to and the share that passed; and the share of turns passing every one."""
    ok = [r["verdicts"] for r in judged.values() if "verdicts" in r]
    summary = {"turns": len(judged), "errors": len(judged) - len(ok), "criteria": {}}
    for name in CRITERIA.values():
        applied = [v[name]["verdict"] for v in ok if v[name]["verdict"] != "n/a"]
        summary["criteria"][name] = {"applies": len(applied), "pass_rate": round(applied.count("pass") / len(applied), 3) if applied else None}
    summary["turn_pass_rate"] = round(sum(all(x["verdict"] != "fail" for x in v.values()) for v in ok) / len(ok), 3) if ok else None
    summary["cost_usd"] = round(sum(r["cost_usd"] for r in judged.values()), 6)
    return summary


def catalog_rows(index: SearchIndex, expected_sha: str) -> dict:
    """Product rows as a search returns them, from the catalog the run used (docs/07-evaluation.md, Metrics and targets)."""
    if catalog_fingerprint(index) != expected_sha:
        raise SystemExit("The catalog doesn't match the one the run used, so product rows can't be rebuilt from IDs.")
    return {row["id"]: _summary(row) for row in index.conn.execute("SELECT * FROM products")}


def results_items(data: dict, rows: dict) -> dict[str, dict]:
    """Every turn of a results file, rendered, keyed `<case>#<turn>`."""
    items = {}
    for case in data["cases"]:
        for i, turn in enumerate(case["turns"]):
            if "tool_calls" not in turn:
                raise SystemExit("This results file has no tool calls; only runs that store them can be judged.")
            items[f"{case['id']}#{i + 1}"] = render({"history": case["turns"][:i], "data": turn}, rows, None)
    return items


async def judge_results_file(path: Path, settings: Settings, index: SearchIndex, holdout: bool = False) -> Path:
    data = json.loads(path.read_text())
    if data.get("eval_set") == "holdout" and not holdout:
        raise SystemExit("Judging holdout needs --holdout, and a plan recorded in evals/README.md first.")
    provider = judge_provider(settings)
    version = judge_version(provider, settings.judge_effort)
    judged = await judge_all(provider, results_items(data, catalog_rows(index, data["catalog_sha256"])))
    out = RESULTS_DIR / data["eval_set"] / "judged" / f"{path.stem}.judge-{version}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    header = {"results_file": path.name, "judge_version": version, "judge_provider": provider.name, "judge_model": provider.model,
              "judge_effort": settings.judge_effort, "summary": summarize_verdicts(judged)}
    out.write_text(json.dumps({**header, "turns": judged}, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps(header["summary"], indent=2), f"Judged: {out}", sep="\n")
    return out


# The labeling sheet: judging its rows, and comparing the judge with the human labels.

def read_sheet(path: Path) -> dict[str, dict]:
    """Each row's information fields and human labels, by row ID."""
    from openpyxl import load_workbook
    sheet = load_workbook(path, read_only=True)["Labels"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    out = {}
    for row in rows:
        if not row or not row[0]:
            continue
        cell = dict(zip(header, row))
        out[row[0]] = {
            "fields": {name: cell[name] for name in INFO},
            "labels": {name: str(cell[name]).strip().lower() if cell.get(name) not in (None, "") else None for name in CRITERIA.values()},
            "notes": cell.get("Notes on fails"),
        }
    return out


def read_key() -> dict:
    return json.loads((LABELS_DIR / "key" / "key.json").read_text())


async def judge_sheet(split: str, settings: Settings) -> Path:
    sheet, key = read_sheet(LABELS_DIR / "labeling-sheet.xlsx"), read_key()
    ids = [r["row"] for r in key["rows"] if r["split"] == split]
    provider = judge_provider(settings)
    version = judge_version(provider, settings.judge_effort)
    judged = await judge_all(provider, {i: sheet[i]["fields"] for i in ids})
    out = LABELS_DIR / "judged" / f"{split}-judge-{version}.json"
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps({"split": split, "judge_version": version, "judge_provider": provider.name, "judge_model": provider.model,
                               "judge_effort": settings.judge_effort, "summary": summarize_verdicts(judged), "rows": judged},
                              indent=2, ensure_ascii=False) + "\n")
    print(f"Judged {len(judged)} {split} rows: {out}")
    return out


def cohen_kappa(pairs: list[tuple[str, str]]) -> float | None:
    """Agreement corrected for chance, over the verdicts both raters gave."""
    if not pairs:
        return None
    observed = sum(a == b for a, b in pairs) / len(pairs)
    first, second = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    expected = sum(first[v] * second[v] for v in VERDICTS) / len(pairs) ** 2
    return 1.0 if expected == 1 else round((observed - expected) / (1 - expected), 3)


def agreement(human: dict[str, dict], other: dict[str, dict]) -> dict:
    """Per criterion: rows compared, agreement, kappa, and the two kinds of error from the human labels' side.

    `human` and `other` map a row ID to verdicts by criterion name; rows or verdicts missing from either are skipped.
    """
    report = {}
    for name in CRITERIA.values():
        pairs = [(human[r][name], other[r][name]) for r in human if r in other and human[r].get(name) in VERDICTS and other[r].get(name) in VERDICTS]
        report[name] = {
            "rows": len(pairs),
            "agreement": round(sum(a == b for a, b in pairs) / len(pairs), 3) if pairs else None,
            "kappa": cohen_kappa(pairs),
            "false_pass": sum(a == "fail" and b == "pass" for a, b in pairs),  # the other rater missed a failure
            "false_fail": sum(a == "pass" and b == "fail" for a, b in pairs),
            "na_mismatch": sum((a == "n/a") != (b == "n/a") for a, b in pairs),
        }
    return report


def print_report(title: str, report: dict) -> None:
    print(f"\n{title}\n{'criterion':24} {'rows':>4} {'agree':>6} {'kappa':>6} {'false pass':>11} {'false fail':>11} {'n/a mismatch':>13}")
    for name, r in report.items():
        agree = f"{r['agreement']:.0%}" if r["agreement"] is not None else "-"
        kappa = f"{r['kappa']:.2f}" if r["kappa"] is not None else "-"
        print(f"{name:24} {r['rows']:>4} {agree:>6} {kappa:>6} {r['false_pass']:>11} {r['false_fail']:>11} {r['na_mismatch']:>13}")


def report_agreement(judged_path: Path) -> None:
    judged, key = json.loads(judged_path.read_text()), read_key()
    sheet = read_sheet(LABELS_DIR / "labeling-sheet.xlsx")
    rows = list(judged["rows"])
    human = {r: sheet[r]["labels"] for r in rows}
    judge = {r: {n: v["verdict"] for n, v in judged["rows"][r]["verdicts"].items()} for r in rows if "verdicts" in judged["rows"][r]}
    unlabeled = [r for r in rows if any(v is None for v in human[r].values())]
    if unlabeled:
        print(f"Rows with a missing human label, skipped for those criteria: {', '.join(unlabeled)}")
    print_report(f"Judge {judged['judge_version']} ({judged['judge_model']}) vs human labels, {judged['split']} rows", agreement(human, judge))
    # The same by the provider that wrote the reply, so a judge favouring one provider's replies shows (docs/07-evaluation.md, Judge).
    provider_of = {k["row"]: k["provider"] for k in key["rows"]}
    for provider in sorted({provider_of[r] for r in rows}):
        mine = [r for r in rows if provider_of[r] == provider]
        print_report(f"... on {provider}'s replies only ({len(mine)} rows)", agreement({r: human[r] for r in mine}, judge))
    broken = [k for k in key["rows"] if k["broken"] and k["row"] in judge]
    caught = [k["row"] for k in broken if judge[k["row"]][k["broken_criterion"]] == "fail"]
    print(f"\nDeliberately broken rows caught (the broken criterion failed): {len(caught)} of {len(broken)}")
    if judged["split"] != "calibration":
        return  # held-out disagreements stay unread, so they can't shape the judge
    print("\nDisagreements (for refining the judge's instructions):")
    for r in rows:
        for name in CRITERIA.values():
            h, j = human[r][name], judge.get(r, {}).get(name)
            if h in VERDICTS and j and h != j:
                print(f"- {r} {name}: human {h}, judge {j}. Judge: {judged['rows'][r]['verdicts'][name]['reason']}"
                      + (f" | Human notes: {sheet[r]['notes']}" if sheet[r]["notes"] else ""))


def report_consistency() -> None:
    key = read_key()
    first = read_sheet(LABELS_DIR / "labeling-sheet.xlsx")
    again = read_sheet(LABELS_DIR / "recheck-sheet.xlsx")
    human = {r: first[source]["labels"] for r, source in key["recheck"].items()}
    print_report("Labeler consistency: recheck sheet vs labeling sheet", agreement(human, {r: again[r]["labels"] for r in key["recheck"]}))


async def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Grade eval turns with the LLM judge, or measure its agreement with human labels.")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--results", type=Path, help="judge every turn of a results file")
    mode.add_argument("--sheet", choices=("calibration", "held-out"), help="judge the labeling sheet's rows of one split")
    mode.add_argument("--agreement", type=Path, help="compare a judged sheet file with the human labels")
    mode.add_argument("--consistency", action="store_true", help="compare the recheck sheet with the labeling sheet")
    parser.add_argument("--holdout", action="store_true", help="allow judging a holdout results file")
    args = parser.parse_args(argv)
    settings = Settings()
    if args.results:
        await judge_results_file(args.results, settings, load_index(settings.data_dir), args.holdout)
    elif args.sheet:
        await judge_sheet(args.sheet, settings)
    elif args.agreement:
        report_agreement(args.agreement)
    else:
        report_consistency()
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
